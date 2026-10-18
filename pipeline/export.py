"""Excel / CSV export.

Spreadsheet safety: any text that starts with ``= + - @`` (or a tab / carriage return) could be run as a
formula by Excel or LibreOffice (CSV/formula injection). All text cells are neutralised by prefixing an
apostrophe, and the .xlsx writer additionally forces text cells to the *string* data type.
"""

from __future__ import annotations

import csv
import io
from datetime import date, datetime, timezone
from decimal import Decimal
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from pipeline.duplicates import DuplicateCandidate
from pipeline.schemas import InvoiceStatus, Severity
from pipeline.storage import InvoiceRow, Storage
from pipeline.validate import normalized_date_or_raw

FORMULA_TRIGGERS = ("=", "+", "-", "@", "\t", "\r", "\n")

INVOICE_COLUMNS = [
    "Source File", "Pages", "Vendor Name", "Vendor GSTIN", "Buyer Name", "Buyer GSTIN", "Invoice Number",
    "Invoice Date", "Taxable Amount", "CGST", "SGST", "IGST", "Total Amount", "Currency",
    "Extraction Status", "Review Status", "Validation Errors", "Validation Warnings",
    "Field Review Flags", "OCR Confidence", "Duplicate Candidate", "Manually Edited Fields", "Model",
]
AMOUNT_COLUMNS = {"Taxable Amount", "CGST", "SGST", "IGST", "Total Amount"}
ISSUE_COLUMNS = ["Source File", "Invoice Number", "Severity", "Code", "Fields", "Explanation"]

REVIEW_LABELS = {
    InvoiceStatus.REVIEWED: "Reviewed",
    InvoiceStatus.NEEDS_REVIEW: "Needs review",
    InvoiceStatus.EXTRACTED: "Not yet reviewed (no flags)",
    InvoiceStatus.FAILED: "Failed",
    InvoiceStatus.PENDING: "Pending",
    InvoiceStatus.PROCESSING: "Processing",
}


def safe_text(value: Any) -> Any:
    """Neutralise spreadsheet formulas in untrusted text. Non-strings are returned unchanged."""
    if isinstance(value, str) and value.startswith(FORMULA_TRIGGERS):
        return "'" + value
    return value


def _dup_text(cands: list[DuplicateCandidate] | None, names: dict[str, str]) -> str:
    if not cands:
        return ""
    return "; ".join(f"{c.strength}: {names.get(c.other_id, c.other_id)} ({c.reason})" for c in cands)


def _flags(row: InvoiceRow) -> str:
    return "; ".join(f"{k}: {v['status']}" for k, v in row.field_status.items() if v.get("status") not in (None, "ok"))


def invoice_records(
    rows: list[InvoiceRow],
    duplicates: dict[str, list[DuplicateCandidate]] | None = None,
    ocr_confidence: dict[str, float] | None = None,
    day_first: bool = True,
    include_unprocessed: bool = False,
) -> list[dict[str, Any]]:
    """One dict per invoice, keyed by INVOICE_COLUMNS. Amounts stay exact ``Decimal`` values (or None)."""
    duplicates = duplicates or {}
    ocr_confidence = ocr_confidence or {}
    names = {r.invoice_id: r.filename for r in rows}
    out: list[dict[str, Any]] = []
    for r in rows:
        if not include_unprocessed and r.status in (InvoiceStatus.PENDING, InvoiceStatus.PROCESSING):
            continue
        inv = r.effective
        errors = [i.short() for i in r.issues if i.severity == Severity.ERROR]
        warnings = [i.short() for i in r.issues if i.severity == Severity.WARNING]
        if r.status == InvoiceStatus.FAILED and r.error:
            errors.insert(0, f"[EXTRACTION_FAILED] {r.error}")
        pages = f"{r.page_start}-{r.page_end}" if r.page_end > r.page_start else str(r.page_start)
        out.append({
            "Source File": r.filename,
            "Pages": pages if r.page_count else "",
            "Vendor Name": inv.vendor_name if inv else None,
            "Vendor GSTIN": inv.vendor_gstin if inv else None,
            "Buyer Name": inv.buyer_name if inv else None,
            "Buyer GSTIN": inv.buyer_gstin if inv else None,
            "Invoice Number": inv.invoice_number if inv else None,
            "Invoice Date": normalized_date_or_raw(inv.invoice_date, day_first) if inv else None,
            "Taxable Amount": inv.taxable_amount if inv else None,
            "CGST": inv.cgst if inv else None,
            "SGST": inv.sgst if inv else None,
            "IGST": inv.igst if inv else None,
            "Total Amount": inv.total_amount if inv else None,
            "Currency": inv.currency if inv else None,
            "Extraction Status": "Failed" if r.status == InvoiceStatus.FAILED else "Succeeded",
            "Review Status": REVIEW_LABELS[r.status],
            "Validation Errors": " | ".join(errors),
            "Validation Warnings": " | ".join(warnings),
            "Field Review Flags": _flags(r),
            "OCR Confidence": ocr_confidence.get(r.invoice_id),
            "Duplicate Candidate": _dup_text(duplicates.get(r.invoice_id), names),
            "Manually Edited Fields": ", ".join(r.edited_fields),
            "Model": r.model or "",
        })
    return out


def issue_records(rows: list[InvoiceRow], include_info: bool = True) -> list[dict[str, Any]]:
    out = []
    for r in rows:
        inv = r.effective
        if r.status == InvoiceStatus.FAILED and r.error:
            out.append({"Source File": r.filename, "Invoice Number": "", "Severity": "error", "Code": "EXTRACTION_FAILED",
                        "Fields": "", "Explanation": r.error})
        for i in r.issues:
            if i.severity == Severity.INFO and not include_info:
                continue
            out.append({"Source File": r.filename, "Invoice Number": (inv.invoice_number if inv else "") or "",
                        "Severity": i.severity.value, "Code": i.code, "Fields": ", ".join(i.fields), "Explanation": i.message})
    return out


# ----------------------------------------------------------------------
def _csv_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, float):
        return f"{value:.4f}"
    return str(safe_text(value))


def to_csv_bytes(records: list[dict[str, Any]], columns: list[str] = INVOICE_COLUMNS) -> bytes:
    """UTF-8 with BOM so Excel shows rupee signs and names correctly."""
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow(columns)
    for rec in records:
        writer.writerow([_csv_cell(rec.get(c)) for c in columns])
    return b"\xef\xbb\xbf" + buffer.getvalue().encode("utf-8")


_HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
_SEVERITY_FILL = {"error": "F8D7DA", "warning": "FFF3CD", "info": "E7F1FF"}


def _write_text_cell(ws: Any, row: int, col: int, value: Any) -> Any:
    cell = ws.cell(row=row, column=col)
    cell.value = safe_text(value)
    if isinstance(cell.value, str):
        cell.data_type = "s"  # never let openpyxl turn text into a formula or error value
    return cell


def _write_table(ws: Any, columns: list[str], records: list[dict[str, Any]], amount_cols: set[str] = frozenset()) -> None:  # type: ignore[assignment]
    for c, name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for r, rec in enumerate(records, start=2):
        for c, name in enumerate(columns, start=1):
            value = rec.get(name)
            if name in amount_cols and isinstance(value, Decimal):
                cell = ws.cell(row=r, column=c, value=value)
                cell.number_format = "#,##0.00"
            elif name == "Invoice Date" and isinstance(value, str) and len(value) == 10 and value[4:5] == "-":
                try:
                    cell = ws.cell(row=r, column=c, value=date.fromisoformat(value))
                    cell.number_format = "yyyy-mm-dd"
                except ValueError:
                    cell = _write_text_cell(ws, r, c, value)
            elif name == "OCR Confidence" and isinstance(value, float):
                cell = ws.cell(row=r, column=c, value=round(value, 4))
                cell.number_format = "0%"
            elif value is None:
                continue
            else:
                cell = _write_text_cell(ws, r, c, value)
            if name in ("Validation Errors", "Validation Warnings", "Field Review Flags", "Explanation", "Duplicate Candidate"):
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        if "Severity" in columns:
            sev = str(rec.get("Severity", ""))
            if sev in _SEVERITY_FILL:
                ws.cell(row=r, column=columns.index("Severity") + 1).fill = PatternFill("solid", fgColor=_SEVERITY_FILL[sev])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(2, len(records) + 1)}"
    for c, name in enumerate(columns, start=1):
        longest = max([len(name)] + [len(str(rec.get(name) or "")) for rec in records[:200]])
        ws.column_dimensions[get_column_letter(c)].width = min(max(10, longest + 2), 60)


def _summary_rows(rows: list[InvoiceRow], records: list[dict[str, Any]]) -> list[tuple[str, Any]]:
    counts: dict[str, int] = {}
    for r in rows:
        counts[r.status.value] = counts.get(r.status.value, 0) + 1
    inr = [x for x in records if (x.get("Currency") or "INR") == "INR"]

    def total(col: str) -> Decimal:
        return sum((x[col] for x in inr if isinstance(x.get(col), Decimal)), Decimal("0"))

    reviewed = sum(1 for x in records if x["Review Status"] == "Reviewed")
    out: list[tuple[str, Any]] = [
        ("Exported at (UTC)", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M")),
        ("Invoices in this export", len(records)),
        ("Reviewed by a person", reviewed),
        ("", ""),
    ]
    out += [(f"Status: {k}", v) for k, v in sorted(counts.items())]
    out += [("", ""), ("PROVISIONAL TOTALS (INR only)", f"Based on {len(inr)} invoice(s); only {reviewed} reviewed. Not for filing until every row is reviewed."),
            ("Taxable amount", total("Taxable Amount")), ("CGST", total("CGST")), ("SGST", total("SGST")),
            ("IGST", total("IGST")), ("Total amount", total("Total Amount"))]
    if len(inr) != len(records):
        out.append(("Note", f"{len(records) - len(inr)} invoice(s) in another currency are excluded from these totals."))
    out.append(("Note", "Amounts are exactly as extracted or corrected. They were never adjusted to make totals add up."))
    return out


def to_xlsx_bytes(
    rows: list[InvoiceRow],
    duplicates: dict[str, list[DuplicateCandidate]] | None = None,
    ocr_confidence: dict[str, float] | None = None,
    day_first: bool = True,
    include_validation_sheet: bool = True,
    include_unprocessed: bool = False,
) -> bytes:
    records = invoice_records(rows, duplicates, ocr_confidence, day_first, include_unprocessed)
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoices"
    _write_table(ws, INVOICE_COLUMNS, records, AMOUNT_COLUMNS)

    if include_validation_sheet:
        issues_ws = wb.create_sheet("Validation")
        issues = issue_records([r for r in rows if include_unprocessed or r.status not in (InvoiceStatus.PENDING, InvoiceStatus.PROCESSING)])
        _write_table(issues_ws, ISSUE_COLUMNS, issues)

    summary = wb.create_sheet("Summary")
    for i, (label, value) in enumerate(_summary_rows(rows, records), start=1):
        _write_text_cell(summary, i, 1, label).font = Font(bold=True)
        if isinstance(value, Decimal):
            cell = summary.cell(row=i, column=2, value=value)
            cell.number_format = "#,##0.00"
        else:
            _write_text_cell(summary, i, 2, value)
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 90

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def validation_report_csv(rows: list[InvoiceRow]) -> bytes:
    return to_csv_bytes(issue_records(rows), ISSUE_COLUMNS)


def gather(storage: Storage, day_first: bool = True) -> tuple[list[InvoiceRow], dict[str, list[DuplicateCandidate]], dict[str, float]]:
    """Load everything an export needs from storage (rows, duplicate candidates, mean OCR confidence)."""
    from pipeline.runner import compute_duplicates  # local import avoids a cycle

    rows = storage.list_invoices()
    confidence: dict[str, float] = {}
    for r in rows:
        values = [p.ocr.mean_confidence for p in storage.get_pages(r.doc_id, r.page_start, r.page_end) if p.ocr.mean_confidence is not None]
        if values:
            confidence[r.invoice_id] = sum(values) / len(values)
    return rows, compute_duplicates(rows, day_first), confidence


def export_all(storage: Storage, fmt: str = "xlsx", day_first: bool = True, include_validation_sheet: bool = True) -> bytes:
    rows, dups, conf = gather(storage, day_first)
    if fmt == "xlsx":
        return to_xlsx_bytes(rows, dups, conf, day_first, include_validation_sheet)
    if fmt == "csv":
        return to_csv_bytes(invoice_records(rows, dups, conf, day_first))
    if fmt == "validation_csv":
        return validation_report_csv(rows)
    raise ValueError(f"Unknown export format '{fmt}'")


def iter_columns() -> Iterable[str]:
    return iter(INVOICE_COLUMNS)
